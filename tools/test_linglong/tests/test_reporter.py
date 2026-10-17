#!/usr/bin/env python3
# SPDX-FileCopyrightText: 2026 Linyaps contributors
#
# SPDX-License-Identifier: LGPL-3.0-or-later

import contextlib
from datetime import datetime, timezone
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from tools.test_linglong import reporter
from tools.test_linglong.models import StepResult


@unittest.skipUnless(reporter.HAS_OPENPYXL, "openpyxl is optional")
class ReporterTextTest(unittest.TestCase):
    def setUp(self):
        self.contexts = contextlib.ExitStack()
        self.addCleanup(self.contexts.close)
        self.directory = self.contexts.enter_context(tempfile.TemporaryDirectory())
        previous = os.getcwd()
        os.chdir(self.directory)
        self.addCleanup(os.chdir, previous)
        self.contexts.enter_context(contextlib.redirect_stdout(io.StringIO()))
        self.json_path = Path(self.directory) / "results.json"

    def write_report(self, text):
        result = StepResult(
            index=7,
            title=text,
            status="FAIL",
            duration_ms=123,
            error_message=text,
            category=text,
        )
        reporter.generate_report(
            [result], datetime(2026, 10, 2, tzinfo=timezone.utc), str(self.json_path)
        )

    def assert_literal_cells(self, expected, data_only=False):
        from openpyxl import load_workbook

        workbook = load_workbook(
            io.BytesIO(Path("玲珑冒烟测试.xlsx").read_bytes()), data_only=data_only
        )
        try:
            for address in ("B4", "C4", "F4"):
                with self.subTest(address=address):
                    cell = workbook.active[address]
                    self.assertEqual(cell.value, expected)
                    self.assertEqual(cell.data_type, "s")
        finally:
            workbook.close()

    def test_formula_like_text_remains_literal(self):
        self.write_report("=1+1")
        self.assert_literal_cells("=1+1")

    def test_formula_like_text_survives_data_only_read(self):
        self.write_report("=1+1")
        self.assert_literal_cells("=1+1", data_only=True)

    def test_saved_workbook_contains_no_formula_elements(self):
        self.write_report("=SUM(1,2)")
        with ZipFile("玲珑冒烟测试.xlsx") as archive:
            sheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        self.assertEqual(sheet.findall(".//s:f", namespace), [])

    def test_excel_error_tokens_remain_literal(self):
        for token in ("#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"):
            with self.subTest(token=token):
                self.write_report(token)
                self.assert_literal_cells(token)

    def test_sanitizing_controls_does_not_create_a_formula(self):
        self.write_report("\x00=1+1\x07")
        self.assert_literal_cells("=1+1")

    def test_regular_unicode_and_newlines_are_preserved(self):
        text = "构建失败\n详细信息\t第二行"
        self.write_report(text)
        self.assert_literal_cells(text)

    def test_numeric_cells_keep_numeric_types_and_status_style(self):
        from openpyxl import load_workbook

        self.write_report("=1+1")
        workbook = load_workbook(io.BytesIO(Path("玲珑冒烟测试.xlsx").read_bytes()))
        try:
            sheet = workbook.active
            for address, value in (("A4", 7), ("E4", 123)):
                self.assertEqual(sheet[address].value, value)
                self.assertEqual(sheet[address].data_type, "n")
            self.assertEqual(sheet["D4"].value, "FAIL")
            self.assertEqual(sheet["D4"].fill.fgColor.rgb, "00FCE4EC")
        finally:
            workbook.close()

    def test_json_preserves_original_text(self):
        text = "\x00=1+1\x07"
        self.write_report(text)
        summary = json.loads(self.json_path.read_text(encoding="utf-8"))
        self.assertEqual(summary["steps"][0]["title"], text)
        self.assertEqual(summary["steps"][0]["error_message"], text)
        self.assertEqual(summary["failed"], 1)


class ReporterWithoutExcelTest(unittest.TestCase):
    def test_json_only_mode_does_not_require_a_workbook(self):
        with (
            tempfile.TemporaryDirectory() as directory,
            patch.object(reporter, "HAS_OPENPYXL", False),
            contextlib.redirect_stdout(io.StringIO()),
            contextlib.redirect_stderr(io.StringIO()),
        ):
            path = Path(directory) / "results.json"
            reporter.generate_report(
                [StepResult(1, "=1+1", "FAIL", 1, "#VALUE!")],
                datetime(2026, 10, 2, tzinfo=timezone.utc),
                str(path),
            )
            summary = json.loads(path.read_text(encoding="utf-8"))
            self.assertEqual(summary["steps"][0]["error_message"], "#VALUE!")


if __name__ == "__main__":
    unittest.main()
