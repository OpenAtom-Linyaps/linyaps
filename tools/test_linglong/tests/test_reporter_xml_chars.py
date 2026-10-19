# SPDX-FileCopyrightText: 2026 huier
#
# SPDX-License-Identifier: LGPL-3.0-or-later

"""Exercise report serialization with characters outside the XML 1.0 range."""

import contextlib
import io
import json
import os
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

from tools.test_linglong.models import StepResult
from tools.test_linglong.reporter import generate_report, sanitize_cell


@unittest.skipUnless(load_workbook is not None, "Excel reporting requires openpyxl")
class ReporterXmlCharactersTest(unittest.TestCase):
    def check_report(self, field, column, text, expected):
        result = StepResult(7, "title", "FAIL", 123, "diagnostic", "module")
        setattr(result, field, text)
        with tempfile.TemporaryDirectory() as directory:
            previous = Path.cwd()
            try:
                os.chdir(directory)
                warnings = io.StringIO()
                with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(warnings):
                    generate_report([result], datetime(2026, 1, 2), "results.json")
                summary = json.loads(Path("results.json").read_text("utf-8"))
                if field in ("title", "error_message"):
                    self.assertEqual(summary["steps"][0][field], text)
                self.assertEqual(summary["failed"], 1)
                self.assertEqual(summary["steps"][0]["duration_ms"], 123)
                self.assertEqual(warnings.getvalue(), "")

                # Check serialized XML as well as an actual workbook round trip.
                path = Path("玲珑冒烟测试.xlsx")
                with ZipFile(path) as archive:
                    for name in archive.namelist():
                        if name.endswith(".xml"):
                            ElementTree.fromstring(archive.read(name))
                workbook = load_workbook(path)
                try:
                    sheet = workbook.active
                    self.assertEqual(sheet.cell(4, column).value, expected)
                    self.assertEqual(sheet["A4"].value, 7)
                    self.assertEqual(sheet["E4"].value, 123)
                    self.assertEqual(sheet["D4"].value, "FAIL")
                    self.assertEqual(sheet["D4"].fill.fgColor.rgb, "00FCE4EC")
                finally:
                    workbook.close()
            finally:
                os.chdir(previous)

    def test_title_noncharacters(self):
        for char in ("\ufffe", "\uffff"):
            with self.subTest(char=ascii(char)):
                self.check_report("title", 3, "before" + char + "after", "beforeafter")

    def test_category_noncharacters(self):
        for char in ("\ufffe", "\uffff"):
            with self.subTest(char=ascii(char)):
                self.check_report("category", 2, "before" + char + "after", "beforeafter")

    def test_error_noncharacters(self):
        for char in ("\ufffe", "\uffff"):
            with self.subTest(char=ascii(char)):
                self.check_report("error_message", 6, "before" + char + "after", "beforeafter")

    def test_valid_xml_characters_survive(self):
        text = "中文\tline\n\ue000\ufffd\U0001fffe\U0010ffff"
        self.assertEqual(sanitize_cell(text), text)
        self.check_report("error_message", 6, text, text)

    def test_existing_control_filter_and_non_strings(self):
        self.assertEqual(sanitize_cell("a\x00\x0b\x1b\x7f\x85b\t\r\n"), "ab\t\r\n")
        self.assertEqual(sanitize_cell(123), 123)
        self.assertIsNone(sanitize_cell(None))


if __name__ == "__main__":
    unittest.main()
